#!/usr/bin/env python3
"""МХБ-ийн биелсэн хэлцлийг өдөр бүр татаж data/deals_YYYY-MM-DD.xlsx файлд хадгална.

Файл бүр: "ALL" хуудас (бүх компани) + компани бүрийн тусдаа хуудас.
Дахин ажиллуулбал хэлцлийн ID-аар давхардлыг арилгаж нэгтгэнэ.
"""
import json
import os
import sys
import time
import urllib.parse
import urllib.request
from collections import defaultdict
from datetime import datetime
from pathlib import Path

from openpyxl import Workbook, load_workbook

API = "http://api.marketinfo.mn/api/mse"
OUT_DIR = Path(os.environ.get("OUT_DIR", "data"))

# Бүртгэх компаниуд (symbol нь <КОД>-O-0000 хэлбэртэй)
CODES = [
    "AARD", "ADB", "APU", "CUMN", "GLMT", "INV", "KHAN", "LEND", "MFC", "MNDL",
    "MSE", "QPAY", "SBM", "TDB", "TTL", "TUM", "XAC", "ERDN", "GAZR", "MFG",
]

HEADERS = ["Огноо цаг", "Symbol", "Үнэ (₮)", "Ширхэг", "Дүн (₮)", "Нөхцөл", "ID"]


def get_json(url, tries=4):
    for i in range(tries):
        try:
            req = urllib.request.Request(url, headers={"User-Agent": "Mozilla/5.0 MSE-collector"})
            with urllib.request.urlopen(req, timeout=30) as r:
                return json.loads(r.read().decode("utf-8"))
        except Exception:
            if i == tries - 1:
                raise
            time.sleep(3 * (i + 1))


def build_code_map():
    """symbol -> тоон код (companycode). /trades-ээс, дутвал /companies-ээс."""
    codes = {}
    try:
        for t in get_json(API + "/trades"):
            if t.get("symbol") and t.get("companycode") is not None:
                codes[t["symbol"]] = t["companycode"]
    except Exception as e:
        print("trades татаж чадсангүй:", e)
    try:
        for c in get_json(API + "/companies"):
            code = c.get("companycode", c.get("code"))
            if c.get("symbol") and code is not None:
                codes.setdefault(c["symbol"], code)
    except Exception as e:
        print("companies татаж чадсангүй:", e)
    return codes


def parse_dt(day, time_str):
    t = (time_str or "").strip()
    if not t or t == "null":
        t = "00:00:00"
    for fmt in ("%Y-%m-%d %H:%M:%S", "%Y-%m-%d %H:%M"):
        try:
            return datetime.strptime(f"{day} {t}", fmt)
        except ValueError:
            pass
    return datetime.strptime(day, "%Y-%m-%d")


def short(sym):
    """APU-O-0000 -> APU, ERDN-O-0001 -> ERDN-1 (өөр хэлбэр бол хэвээр)."""
    p = str(sym).split("-")
    if len(p) == 3 and p[1] == "O" and p[2].isdigit():
        return p[0] if int(p[2]) == 0 else f"{p[0]}-{int(p[2])}"
    return str(sym)


def to_row(r):
    day = str(r.get("dates", ""))[:10]
    px = float(r["mdEntryPx"])
    sz = float(r["mdEntrySize"])
    cond = r.get("tradeCondition")
    cond = "" if cond in (None, "null") else cond
    row = [parse_dt(day, r.get("mdEntryTime")), short(r["symbol"]), px, int(sz) if sz == int(sz) else sz,
           px * sz, cond, r["id"]]
    return day, row


def read_existing(path):
    if not path.exists():
        return {}
    ws = load_workbook(path, read_only=True)["ALL"]
    rows = {}
    for i, r in enumerate(ws.iter_rows(values_only=True)):
        if i == 0 or r[6] is None:
            continue
        row = list(r[:7])
        row[1] = short(row[1])  # хуучин файлын урт кодыг богиносгоно
        rows[r[6]] = row
    return rows


def fill(ws, rows):
    ws.append(HEADERS)
    for r in rows:
        ws.append(r)
    for i, w in enumerate([20, 14, 14, 12, 16, 20, 12], start=1):
        ws.column_dimensions[chr(64 + i)].width = w
    for row in ws.iter_rows(min_row=2):
        row[0].number_format = "yyyy-mm-dd hh:mm:ss"
        row[2].number_format = "#,##0.00"
        row[3].number_format = "#,##0"
        row[4].number_format = "#,##0.00"
    ws.freeze_panes = "A2"


def write_xlsx(path, rows_by_id):
    rows = sorted(rows_by_id.values(), key=lambda r: (r[0], r[6]))
    wb = Workbook()
    ws = wb.active
    ws.title = "ALL"
    fill(ws, rows)
    by_sym = defaultdict(list)
    for r in rows:
        by_sym[r[1]].append(r)
    for sym in sorted(by_sym):
        fill(wb.create_sheet(sym[:31]), by_sym[sym])
    OUT_DIR.mkdir(parents=True, exist_ok=True)
    wb.save(path)


def main():
    code_map = build_code_map()
    targets, missing = [], []
    for code in CODES:
        syms = [s for s in code_map if s.startswith(code + "-O-")]
        if syms:
            targets += [(s, code_map[s]) for s in sorted(syms)]
        else:
            missing.append(code)
    if missing:
        print("Тоон код олдсонгүй (алгасав):", ", ".join(missing))

    by_day = defaultdict(dict)  # өдөр -> {id: мөр}
    failed = []
    for sym, num in targets:
        try:
            rows = get_json(f"{API}/deals/{num}/{urllib.parse.quote(sym)}")
            for r in rows if isinstance(rows, list) else []:
                day, row = to_row(r)
                by_day[day][row[6]] = row
            print(f"{sym}: {len(rows) if isinstance(rows, list) else 0} хэлцэл")
        except Exception as e:
            failed.append(sym)
            print(f"{sym}: АЛДАА {e}")

    for day, new_rows in by_day.items():
        path = OUT_DIR / f"deals_{day}.xlsx"
        merged = read_existing(path)
        merged.update(new_rows)
        write_xlsx(path, merged)
        print(f"Хадгаллаа: {path} ({len(merged)} мөр)")

    if not by_day:
        print("Хадгалах хэлцэл олдсонгүй.")
    if failed:
        print("Амжилтгүй:", ", ".join(failed))
        return 1
    return 0


if __name__ == "__main__":
    sys.exit(main())
