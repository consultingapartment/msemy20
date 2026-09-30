#!/usr/bin/env python3
"""МХБ-ийн хэлцлийн шинжилгээ.

    python analyze.py daily  [YYYY-MM-DD]   # өдрийн шинжилгээ (анхдагч: сүүлийн өгөгдөлтэй өдөр)
    python analyze.py weekly [YYYY-MM-DD]   # тухайн өдөр агуулсан 7 хоногийн (Да–Ба) шинжилгээ
    python analyze.py weekly-auto           # автомат: Баасан гарагт энэ 7 хоног, бусад өдөр
                                            # өмнөх 7 хоногийн тайлан дутуу бол нөхөж гаргана

Оролт : data/deals_YYYY-MM-DD.xlsx   (collect.py үүсгэсэн)
Гаралт: reports/daily_*.xlsx|txt, reports/weekly_*.xlsx|txt, reports/latest_daily.txt, latest_weekly.txt

Үнийн тодорхойлолт:
  Хаалт            = тухайн өдрийн (10:00–13:00) НИЙТ хэлцлийн жигнэсэн дундаж ханш (Σ үнэ×ширхэг / Σ ширхэг)
  VWAP             = тайлангийн бүх хугацааны нийт арилжааны жигнэсэн дундаж (өдрийн тайланд хаалттай тэнцүү)
  Албан ёсны хаалт = сүүлийн 1 цагийн хэлцлийн жигнэсэн дундаж (МХБ-ийн нийтэлдэг), лавлагаа болгон хадгална
"""
import os
import re
import shutil
import sys
from datetime import date, datetime, timedelta, timezone
from pathlib import Path

import pandas as pd
from openpyxl.styles import Font
from openpyxl.utils import get_column_letter

DATA_DIR = Path(os.environ.get("DATA_DIR", "data"))
REPORT_DIR = Path(os.environ.get("REPORT_DIR", "reports"))
CLOSE_TIME = os.environ.get("CLOSE_TIME", "13:00")                  # арилжааны төгсгөл
BLOCK_VALUE = float(os.environ.get("BLOCK_VALUE", 10_000_000))      # том хэлцлийн босго (₮)
JUMP_PCT = float(os.environ.get("JUMP_PCT", 2.0))                   # үнийн үсрэлтийн босго (%)
ORDER_SIZE = float(os.environ.get("ORDER_SIZE", 5_000_000))         # хөрвөх чадварын жишиг захиалга (₮)
THIN_PCT = float(os.environ.get("THIN_PCT", 50))                    # захиалга эргэлтийн энэ %-аас дээш бол "сул"
SINGLE_SHARE = float(os.environ.get("SINGLE_SHARE", 25))            # нэг хэлцлийн эзлэх хувийн босго (%)
SINGLE_MOVE = float(os.environ.get("SINGLE_MOVE", 2.0))             # үнийн өөрчлөлтийн босго (%)
CLOSE_GAP = float(os.environ.get("CLOSE_GAP", 1.0))                 # хаалт-VWAP зөрүүний босго (%)
MIN_DEALS = 5
MAX_CHARS = int(os.environ.get("MAX_CHARS", 3500))                   # тайлангийн дээд урт (Telegram-ийн хязгаар 4096)
MIN_LIQ = float(os.environ.get("MIN_LIQ", 1_000_000))                # хөрвөх чадварын өөрчлөлтөд орох доод дүн (₮)

WEEKDAYS = ["Даваа", "Мягмар", "Лхагва", "Пүрэв", "Баасан", "Бямба", "Ням"]
SHORT_WD = ["Да", "Мя", "Лх", "Пү", "Ба", "Бя", "Ня"]
DISCLAIMER = "⚠️ Энэ нь хөрөнгө оруулалтын зөвлөгөө биш, зөвхөн биелсэн хэлцлийн статистик тойм юм."
RENAME = {"Огноо цаг": "dt", "Symbol": "sym", "Үнэ (₮)": "px", "Ширхэг": "qty",
          "Дүн (₮)": "val", "Нөхцөл": "cond", "ID": "id"}
COL_IMPACT = f"{ORDER_SIZE / 1e6:g} сая ₮ захиалга / эргэлт (%)"


# ---------------------------------------------------------------- өгөгдөл унших
def all_days():
    days = []
    for p in DATA_DIR.glob("deals_*.xlsx"):
        m = re.match(r"deals_(\d{4}-\d{2}-\d{2})\.xlsx$", p.name)
        if m:
            days.append(m.group(1))
    return sorted(days)


def load_day(d):
    df = pd.read_excel(DATA_DIR / f"deals_{d}.xlsx", sheet_name="ALL").rename(columns=RENAME)
    df["dt"] = pd.to_datetime(df["dt"])
    df["cond"] = df["cond"].fillna("").astype(str)
    df["day"] = df["dt"].dt.strftime("%Y-%m-%d")
    return df.sort_values(["dt", "id"]).reset_index(drop=True)


def official_close(g):
    """Хаалтын ханш: сүүлийн арилжааны өдрийн сүүлийн 1 цагийн жигнэсэн дундаж ханш.
    Сүүлийн 1 цагт хэлцэл байхгүй бол сүүлийн хэлцлийн үнийг авч, суурийг тэмдэглэнэ."""
    gd = g[g["day"] == g["day"].max()]
    end = pd.Timestamp(f"{gd['day'].iloc[0]} {CLOSE_TIME}")
    w = gd[(gd["dt"] > end - pd.Timedelta(hours=1)) & (gd["dt"] <= end)]
    if len(w) and w["qty"].sum() > 0:
        return round(w["val"].sum() / w["qty"].sum(), 2), "сүүлийн 1 цагийн жигнэсэн дундаж"
    return float(gd.sort_values(["dt", "id"])["px"].iloc[-1]), "сүүлийн хэлцэл (сүүлийн 1 цагт хэлцэлгүй)"


def day_stats(d):
    """Өмнөх өдрийн суурь: хаалт = тэр өдрийн нийт хэлцлийн жигнэсэн дундаж ханш."""
    return {sym: {"close": round(g["val"].sum() / g["qty"].sum(), 2)}
            for sym, g in load_day(d).groupby("sym")}


# ---------------------------------------------------------------- тооцоолол
def executions(df):
    """Нэг компанийн, яг ижил секундэд биелсэн хэлцлүүдийг нэг хэлцэл болгон нэгтгэнэ.
    Шалтгаан: нэг том захиалга хэд хэдэн жижиг захиалгатай таарахад олон мөр үүсдэг.
    Нэгтгэсэн үнэ = жигнэсэн дундаж (Σ үнэ×ширхэг / Σ ширхэг)."""
    e = df.groupby(["sym", "dt"], sort=False).agg(
        n=("id", "count"), qty=("qty", "sum"), val=("val", "sum"),
        pmin=("px", "min"), pmax=("px", "max"),
        cond=("cond", lambda c: ", ".join(sorted({x for x in c if x})))).reset_index()
    e["px"] = e["val"] / e["qty"]
    return e


def summarize(df, prev=None):
    total = df["val"].sum()
    rows = []
    for sym, g in df.groupby("sym"):
        g = g.sort_values(["dt", "id"])
        vol, val = g["qty"].sum(), g["val"].sum()
        vwap = val / vol
        gd = g[g["day"] == g["day"].max()]                       # сүүлийн арилжааны өдөр
        close = round(gd["val"].sum() / gd["qty"].sum(), 2)      # хаалт = өдрийн нийт жигнэсэн дундаж
        off, obasis = official_close(g)                          # албан ёсны (сүүлийн 1 цаг)
        o, last = g["px"].iloc[0], g["px"].iloc[-1]
        hi, lo = g["px"].max(), g["px"].min()
        p = (prev or {}).get(sym, {})
        pc, pv = p.get("close"), p.get("vwap")
        base, basis = (pc, "өмнөх хаалт") if pc else (o, "нээлт")
        e = executions(g)
        big = e.loc[e["val"].idxmax()]
        pq = g.groupby("px")["qty"].sum()
        rows.append({
            "Symbol": sym,
            "Хэлцэл": len(g),
            "Ширхэг": int(vol),
            "Эргэлт (₮)": float(val),
            "Эзлэх хувь (%)": val / total * 100 if total else 0.0,
            "Нээлт": o,
            "Сүүлийн хэлцэл": last,
            "Хаалт": close,
            "VWAP": vwap,
            "Албан ёсны хаалт": off,
            "Албан ёсны хаалтын суурь": obasis,
            "Сүүлийн цаг − хаалт (%)": (off / close - 1) * 100,
            "Дээд": hi, "Доод": lo,
            "Хамгийн их арилжсан үнэ": float(pq.idxmax()),
            "Тэр үнийн ширхэгийн эзлэх (%)": float(pq.max() / vol * 100),
            "Өмнөх хаалт": pc,
            "Өмнөх VWAP": pv,
            "Өөрчлөлт (%)": (close / base - 1) * 100 if base else 0.0,
            "Өөрчлөлтийн суурь": basis,
            "VWAP өөрчлөлт (%)": (vwap / pv - 1) * 100 if pv else None,
            "Хэлбэлзэл (%)": (hi / lo - 1) * 100 if lo else 0.0,
            "Дундаж хэлцэл (ш)": vol / len(g),
            "Медиан хэлцэл (ш)": float(g["qty"].median()),
            "Хамгийн том хэлцэл (ш)": int(big["qty"]),
            "Хамгийн том хэлцлийн нэгтгэсэн тоо": int(big["n"]),
            "Хамгийн том хэлцлийн эзлэх (%)": float(big["val"]) / val * 100,
            "Том хэлцэл (тоо)": int((e["val"] >= BLOCK_VALUE).sum()),
            COL_IMPACT: ORDER_SIZE / val * 100,
        })
    out = pd.DataFrame(rows).sort_values("Эргэлт (₮)", ascending=False).reset_index(drop=True)
    if out["Өмнөх VWAP"].isna().all():  # өдрийн тайланд өмнөх VWAP хэрэггүй (хаалттай адил)
        out = out.drop(columns=["Өмнөх VWAP", "VWAP өөрчлөлт (%)"])
    return out


def definitions():
    return pd.DataFrame({"Нэр томьёо": ["Хаалт", "VWAP", "Албан ёсны хаалт", "Өөрчлөлт (%)", "Сүүлийн цаг − хаалт (%)",
                                        COL_IMPACT, "Нэгтгэсэн хэлцэл", "Хөрвөх чадварын өөрчлөлт"],
                         "Тодорхойлолт": [
        "Сүүлийн арилжааны өдрийн (10:00–13:00) НИЙТ хэлцлийн жигнэсэн дундаж ханш = Σ(үнэ×ширхэг) / Σ ширхэг",
        "Тайлангийн бүх хугацааны нийт арилжааны жигнэсэн дундаж ханш (өдрийн тайланд хаалттай тэнцүү)",
        "Сүүлийн 1 цагийн (12:00–13:00) хэлцлийн жигнэсэн дундаж. МХБ-ийн нийтэлдэг хаалттай харьцуулах лавлагаа",
        "Хаалтыг өмнөх арилжааны өдрийн хаалттай харьцуулсан. Өмнөх өдөр байхгүй бол эхний хэлцлийн үнэтэй",
        "Албан ёсны хаалт хаалтаас хэдэн хувиар өндөр/доогуур байна. Төгсгөлийн худалдан авалт, зарлагын хүч",
        "Жишиг захиалга нь өдрийн нийт эргэлтийн хэдэн хувь вэ. Их байх тусам хөрвөх чадвар сул",
        "Нэг компанийн, яг ижил секундэд биелсэн хэлцлүүдийг нэг хэлцэл гэж тооцсон. Нэг том захиалга хэд хэдэн "
        "жижиг захиалгатай таарахад олон мөр үүсдэг. Нэгтгэсэн үнэ = жигнэсэн дундаж, ширхэг ба дүн = нийлбэр",
        "Компанийн арилжааны дүнг өмнөх өдөртэй (7 хоногийн тайланд өдрийн дундажаар өмнөх 7 хоногтой) харьцуулсан. "
        "Нэмэгдэх нь хөрвөх чадвар сайжирсан, буурах нь муудсан гэсэн үг"]})


def block_table(df, n=50):
    e = executions(df)
    b = e[e["val"] >= BLOCK_VALUE].sort_values("val", ascending=False).head(n)
    return b[["dt", "sym", "n", "qty", "px", "val", "cond"]].rename(columns={
        "dt": "Огноо цаг", "sym": "Symbol", "n": "Нэгтгэсэн хэлцэл (тоо)", "qty": "Ширхэг",
        "px": "Үнэ (жигнэсэн дундаж)", "val": "Дүн (₮)", "cond": "Нөхцөл"})


def biggest_by_company(df):
    """Компани бүрийн хамгийн том (нэгтгэсэн) хэлцэл ба түүний эзлэх хувь."""
    e = executions(df)
    total = df["val"].sum()
    ct = df.groupby("sym")["val"].sum()
    b = e.loc[e.groupby("sym")["val"].idxmax()].copy()
    b["mkt"] = b["val"] / total * 100
    b["co"] = [v / ct[s_] * 100 for v, s_ in zip(b["val"], b["sym"])]
    return b.sort_values("val", ascending=False).reset_index(drop=True)


def biggest_table(df):
    b = biggest_by_company(df)
    return b[["sym", "dt", "n", "qty", "px", "pmin", "pmax", "val", "mkt", "co", "cond"]].rename(columns={
        "sym": "Symbol", "dt": "Огноо цаг", "n": "Нэгтгэсэн хэлцэл (тоо)", "qty": "Ширхэг",
        "px": "Үнэ (жигнэсэн дундаж)", "pmin": "Доод үнэ", "pmax": "Дээд үнэ", "val": "Дүн (₮)",
        "mkt": "Нийт арилжаанд эзлэх (%)", "co": "Компанийн эргэлтэд эзлэх (%)", "cond": "Нөхцөл"})


def jump_table(df, n=100):
    d = df.sort_values(["sym", "dt", "id"]).copy()
    d["prev"] = d.groupby(["sym", "day"])["px"].shift(1)
    d["chg"] = (d["px"] / d["prev"] - 1) * 100
    j = d[d["chg"].abs() >= JUMP_PCT].copy()
    j["_a"] = j["chg"].abs()
    j = j.sort_values("_a", ascending=False).head(n)
    return j[["dt", "sym", "prev", "px", "chg", "qty"]].rename(columns={
        "dt": "Огноо цаг", "sym": "Symbol", "prev": "Өмнөх хэлцлийн үнэ", "px": "Үнэ (₮)",
        "chg": "Өөрчлөлт (%)", "qty": "Ширхэг"})


def hourly(df):
    h = df.assign(h=df["dt"].dt.hour).groupby("h").agg(
        deals=("id", "count"), qty=("qty", "sum"), val=("val", "sum")).reset_index()
    h["Цаг"] = h["h"].map(lambda x: f"{x:02d}:00–{x:02d}:59")
    h["Эзлэх хувь (%)"] = h["val"] / h["val"].sum() * 100
    return h[["Цаг", "deals", "qty", "val", "Эзлэх хувь (%)"]].rename(columns={
        "deals": "Хэлцэл", "qty": "Ширхэг", "val": "Эргэлт (₮)"})


def condition_table(df):
    c = df[df["cond"] != ""]
    return c[["dt", "sym", "px", "qty", "val", "cond"]].rename(columns={
        "dt": "Огноо цаг", "sym": "Symbol", "px": "Үнэ (₮)", "qty": "Ширхэг", "val": "Дүн (₮)", "cond": "Нөхцөл"})


# ---------------------------------------------------------------- Excel
def save_book(path, sheets):
    path.parent.mkdir(parents=True, exist_ok=True)
    with pd.ExcelWriter(path, engine="openpyxl") as xw:
        for name, df in sheets.items():
            name = name[:31]
            df.to_excel(xw, sheet_name=name, index=False)
            ws = xw.sheets[name]
            ws.freeze_panes = "A2"
            for cell in ws[1]:
                cell.font = Font(bold=True)
            for i, col in enumerate(df.columns, start=1):
                letter = get_column_letter(i)
                series = df[col]
                if pd.api.types.is_datetime64_any_dtype(series):
                    fmt, width = "yyyy-mm-dd hh:mm:ss", 20
                elif pd.api.types.is_float_dtype(series):
                    fmt = "#,##0.00"
                    width = max([len(str(col))] + [len(f"{v:,.2f}") for v in series.head(200)]) + 2
                elif pd.api.types.is_integer_dtype(series):
                    fmt = "#,##0"
                    width = max([len(str(col))] + [len(f"{v:,}") for v in series.head(200)]) + 2
                else:
                    fmt = None
                    width = max([len(str(col))] + [len(str(v)) for v in series.head(200)]) + 2
                ws.column_dimensions[letter].width = min(max(width, 8), 30)
                if fmt:
                    for cell in ws[letter][1:]:
                        cell.number_format = fmt


# ---------------------------------------------------------------- тайлангийн текст (Telegram)
def mn(x):
    if x >= 1e9:
        return f"{x / 1e9:,.2f} тэрбум ₮"
    if x >= 1e6:
        return f"{x / 1e6:,.1f} сая ₮"
    return f"{x:,.0f} ₮"


def pct(x):
    return f"{x:+.2f}%"


def px(x):
    return f"{x:,.0f}" if x >= 1000 else f"{x:,.2f}".rstrip("0").rstrip(".")


def sm(x):
    return f"{x / 1e6:.1f} сая" if x >= 1e6 else f"{x / 1e3:.0f} мянга"


def wd(day):
    return WEEKDAYS[date.fromisoformat(day).weekday()]


def basis_text(s, weekly=False):
    b = set(s["Өөрчлөлтийн суурь"])
    if b == {"өмнөх хаалт"}:
        return "өмнөх хаалттай"
    if b == {"нээлт"}:
        return "нээлтийн үнэтэй"
    return "өмнөх хаалт/нээлттэй"


def liq_rows(cur, base):
    """Хөрвөх чадварын өөрчлөлт: одоогийн арилжааны дүн өмнөхтэй харьцуулахад."""
    rows = []
    for sym, c in cur.items():
        b = base.get(sym, 0.0)
        if b <= 0 or max(c, b) < MIN_LIQ:
            continue
        rows.append((sym, b, c, c / b - 1))
    return rows


def liq_table(cur, base):
    rows = sorted(liq_rows(cur, base), key=lambda r: -r[3])
    return pd.DataFrame(rows, columns=["Symbol", "Өмнөх арилжааны дүн (₮)", "Одоогийн арилжааны дүн (₮)", "_r"]).assign(
        **{"Өөрчлөлт (%)": lambda d: d["_r"] * 100, "Дахин": lambda d: d["Одоогийн арилжааны дүн (₮)"] / d["Өмнөх арилжааны дүн (₮)"]}
    ).drop(columns="_r")


# ---- блокууд: мөр бүхий жагсаалт буцаана
def index_block(s, weekly=False):
    w, ch = s["Эргэлт (₮)"], s["Өөрчлөлт (%)"]
    eq, vw = ch.mean(), (ch * w).sum() / w.sum()
    up, dn = int((ch > 0.05).sum()), int((ch < -0.05).sum())
    contrib = ch * w / w.sum()
    lo, hi = contrib.idxmin(), contrib.idxmax()
    L = [f"📈 МХБ-н 20{' (7 хоног)' if weekly else ''}: тэнцүү {pct(eq)} · жигнэсэн {pct(vw)} · "
         f"↑{up} ↓{dn} ={len(s) - up - dn} ({basis_text(s, weekly)})"]
    parts = []
    if contrib[lo] < 0:
        parts.append(f"татсан {s.loc[lo, 'Symbol']} {contrib[lo]:+.2f} п.п.")
    if contrib[hi] > 0:
        parts.append(f"түлхсэн {s.loc[hi, 'Symbol']} {contrib[hi]:+.2f} п.п.")
    parts.append(f"эхний 4 компани {s.head(4)['Эзлэх хувь (%)'].sum():.0f}%")
    txt = " · ".join(parts)
    L.append(txt[0].upper() + txt[1:])
    return L


def price_block(s, weekly=False, limit=None):
    rows = s if limit is None else s.head(limit)
    if weekly:
        L = ["💱 7 хоногийн жигнэсэн дундаж → хаалт (өөрчлөлт) · мужийн өргөн % · арилжааны дүн:"]
        L += [f"{r['Symbol']} {px(r['VWAP'])} → {px(r['Хаалт'])} ({pct(r['Өөрчлөлт (%)'])}) · {r['Хэлбэлзэл (%)']:.1f}% · {sm(r['Эргэлт (₮)'])}"
              for _, r in rows.iterrows()]
    else:
        L = ["💱 Хаалт (өдрийн жигнэсэн дундаж ханш, өөрчлөлт) · арилжааны дүн · хэлцэл:"]
        L += [f"{r['Symbol']} {px(r['Хаалт'])} ({pct(r['Өөрчлөлт (%)'])}) · {sm(r['Эргэлт (₮)'])} · {int(r['Хэлцэл'])} хэлцэл"
              for _, r in rows.iterrows()]
    return L


def movers_block(s, k=3):
    t = s[s["Хэлцэл"] >= MIN_DEALS]
    up = t[t["Өөрчлөлт (%)"] > 0].sort_values("Өөрчлөлт (%)", ascending=False).head(k)
    dn = t[t["Өөрчлөлт (%)"] < 0].sort_values("Өөрчлөлт (%)").head(k)
    L = []
    if len(up):
        L.append("📈 Өсөлт: " + ", ".join(f"{r['Symbol']} {pct(r['Өөрчлөлт (%)'])}" for _, r in up.iterrows()))
    if len(dn):
        L.append("📉 Уналт: " + ", ".join(f"{r['Symbol']} {pct(r['Өөрчлөлт (%)'])}" for _, r in dn.iterrows()))
    return L


def close_block(s):
    t = s[s["Хэлцэл"] >= MIN_DEALS]
    col = "Сүүлийн цаг − хаалт (%)"
    weak = t[t[col] <= -CLOSE_GAP].sort_values(col).head(3)
    strong = t[t[col] >= CLOSE_GAP].sort_values(col, ascending=False).head(3)
    L = []
    if len(weak):
        L.append("🔻 Төгсгөлд суларсан (сүүлийн 1 цагийн ханш хаалтаас доор): " + ", ".join(
            f"{r['Symbol']} {pct(r[col])}" for _, r in weak.iterrows()))
    if len(strong):
        L.append("🔺 Төгсгөлд хүчирхэгжсэн (дээгүүр): " + ", ".join(
            f"{r['Symbol']} {pct(r[col])}" for _, r in strong.iterrows()))
    t2 = s[(s["Өөрчлөлт (%)"].abs() >= SINGLE_MOVE) & (s["Хамгийн том хэлцлийн эзлэх (%)"] >= SINGLE_SHARE)]
    for _, r in t2.sort_values("Өөрчлөлт (%)", key=lambda x: -x.abs()).iterrows():
        mg = f", {int(r['Хамгийн том хэлцлийн нэгтгэсэн тоо'])} хэлцэл нэгтгэсэн" if r["Хамгийн том хэлцлийн нэгтгэсэн тоо"] > 1 else ""
        L.append(f"⚠️ Нэг хэлцлээр хөдөлсөн байж болзошгүй: {r['Symbol']} {pct(r['Өөрчлөлт (%)'])} — эргэлтийн "
                 f"{r['Хамгийн том хэлцлийн эзлэх (%)']:.0f}% нь ганц хэлцэл ({int(r['Хамгийн том хэлцэл (ш)']):,} ш{mg})")
    return L


def pin_block(s, k=4):
    t = s[(s["Хэлцэл"] >= MIN_DEALS) & (s["Тэр үнийн ширхэгийн эзлэх (%)"] >= 50)]
    t = t.sort_values("Тэр үнийн ширхэгийн эзлэх (%)", ascending=False).head(k)
    if t.empty:
        return []
    return ["🎯 Нэг үнэд бэхлэгдсэн (тэр үнээр арилжсан ширхэгийн хувь): " + ", ".join(
        f"{r['Symbol']} {px(r['Хамгийн их арилжсан үнэ'])} ({r['Тэр үнийн ширхэгийн эзлэх (%)']:.0f}%)" for _, r in t.iterrows())]


def hours_block(df, with_open=True):
    h = hourly(df)
    L = ["⏱ Арилжааны дүн цагаар: " + " · ".join(
        f"{r['Цаг'][:2]}ц {r['Эзлэх хувь (%)']:.0f}%" for _, r in h.iterrows())]
    if with_open:
        total = df["val"].sum()
        day = df["day"].iloc[0]
        o10 = df[df["dt"] < pd.Timestamp(f"{day} 10:10")]["val"].sum() / total * 100
        si = df[df["cond"].str.contains("Sell Imbalance")]["val"].sum() / total * 100
        L[0] += f" | нээлтийн эхний 10 мин {o10:.0f}% (Sell Imbalance хэлцэл {si:.0f}%)"
    return L


def jumps_block(df):
    j = jump_table(df, n=100000)
    if j.empty:
        return []
    big = j.iloc[0]
    cnt = j["Symbol"].value_counts()
    return [f"⚡ Үнийн үсрэлт (дараалсан хоёр хэлцэл ≥{JUMP_PCT:g}%): {len(j)} удаа · хамгийн том {big['Symbol']} "
            f"{big['Огноо цаг']:%H:%M:%S} {big['Өөрчлөлт (%)']:+.1f}% ({px(big['Өмнөх хэлцлийн үнэ'])}→{px(big['Үнэ (₮)'])}, "
            f"{int(big['Ширхэг']):,} ш) · хамгийн олон: {cnt.index[0]} ({cnt.iloc[0]})"]


def biggest_block(df, weekly=False, limit=None):
    b = biggest_by_company(df)
    if limit:
        b = b.head(limit)
    L = ["🐋 Компани бүрийн хамгийн том хэлцэл (цаг · дүн · зах зээлд % · компанид %):"]
    for i, r in b.iterrows():
        when = f"{SHORT_WD[r['dt'].weekday()]} {r['dt']:%H:%M:%S}" if weekly else f"{r['dt']:%H:%M:%S}"
        m = f" ×{int(r['n'])}" if r["n"] > 1 else ""
        c = " SI" if "Sell Imbalance" in r["cond"] else ""
        L.append(f"{i + 1}. {r['sym']} {when} · {sm(r['val'])} · {r['mkt']:.2f}% · {r['co']:.0f}%{m}{c}")
    L.append("×N = нэг секундэд нэг компанийн N хэлцлийг нэгтгэсэн (нэг захиалга байж болзошгүй) · SI = Sell Imbalance")
    return L


def liquidity_block(s, span, change=None, k=3):
    L = ["💧 Хөрвөх чадвар"]
    if change:
        cur, base, label = change
        rows = liq_rows(cur, base)
        up = sorted([r for r in rows if r[3] > 0], key=lambda r: -r[3])[:k]
        dn = sorted([r for r in rows if r[3] < 0], key=lambda r: r[3])[:k]

        def f(r):
            ratio = r[2] / r[1]
            return f"{r[0]} {r[1] / 1e6:.1f}→{r[2] / 1e6:.1f} ({f'{ratio:.1f}×' if ratio >= 2 else f'{r[3] * 100:+.0f}%'})"

        if up or dn:
            L[0] += f" (арилжааны дүн, сая ₮; {label})"
        if up:
            L.append("📈 Нэмэгдсэн: " + " · ".join(f(r) for r in up))
        if dn:
            L.append("📉 Буурсан: " + " · ".join(f(r) for r in dn))
    best = s.sort_values(COL_IMPACT).head(3)
    thin = s[s[COL_IMPACT] >= THIN_PCT].sort_values(COL_IMPACT, ascending=False).head(6)
    line = f"{ORDER_SIZE / 1e6:g} сая ₮ захиалга / {span} арилжааны дүн: сайн " + ", ".join(
        f"{r['Symbol']} {r[COL_IMPACT]:.0f}%" for _, r in best.iterrows())
    if len(thin):
        line += f" · сул ({THIN_PCT:.0f}%+) " + ", ".join(f"{r['Symbol']} {r[COL_IMPACT]:.0f}%" for _, r in thin.iterrows())
    L.append(line)
    return L


def assemble(blocks, limit=None):
    """Блокуудыг нэг текст болгоно; MAX_CHARS-аас урт бол дараах дарааллаар багасгана:
    1) p>=3 блокуудыг (ач холбогдол бага) хамгийн бага ач холбогдлоос нь хасна
    2) компанийн жагсаалтуудыг (fn) нэг нэг мөрөөр багасгана (доод хязгаар nmin)
    3) p=2 блокуудыг сүүлээс нь хасна. p<=1 блок хэзээ ч хасагдахгүй."""
    limit = limit or MAX_CHARS
    bs = [dict(b) for b in blocks]

    def lines_of(b):
        return b["fn"](b["n"]) if b.get("fn") else b["lines"]

    def render():
        return "\n\n".join("\n".join(lines_of(b)) for b in bs if lines_of(b))

    def remove(b):
        b["fn"], b["lines"] = None, []

    while len(render()) > limit:
        c3 = [b for b in bs if b["p"] >= 3 and lines_of(b)]
        if c3:
            remove(max(c3, key=lambda x: x["p"]))
            continue
        cv = [b for b in bs if b.get("fn") and b["n"] > b["nmin"]]
        if cv:
            b = max(cv, key=lambda x: len("\n".join(lines_of(x))))
            b["n"] -= 1
            continue
        c2 = [b for b in bs if b["p"] == 2 and b.get("drop", True) and lines_of(b)]
        if c2:
            remove(c2[-1])
            continue
        break
    return render()


def blk(p, lines=None, drop=True, fn=None, n=None, nmin=12):
    return {"p": p, "lines": lines or [], "drop": drop, "fn": fn, "n": n, "nmin": nmin}


def daily_text(day, df, s, liq=None, prev=None):
    total = df["val"].sum()
    head = [f"📊 МХБ {day} ({wd(day)}) · {len(df):,} хэлцэл · {mn(total)} · {len(s)} компани"]
    if prev:
        pday, pval, pdeals = prev
        head.append(f"Өмнөх өдөртэй ({pday[5:]}) харьцуулахад арилжааны дүн {pct((total / pval - 1) * 100)} "
                    f"({mn(pval)} → {mn(total)}), хэлцэл {pdeals:,} → {len(df):,}")
    blocks = [
        blk(0, head),
        blk(1, index_block(s)),
        blk(2, fn=lambda n: price_block(s, limit=n), n=len(s), drop=False),
        blk(2, movers_block(s)),
        blk(2, close_block(s)),
        blk(3, pin_block(s)),
        blk(3, hours_block(df)),
        blk(4, jumps_block(df)),
        blk(2, fn=lambda n: biggest_block(df, limit=n), n=len(s), drop=False),
        blk(1, liquidity_block(s, "өдрийн", liq)),
        blk(0, [DISCLAIMER]),
    ]
    return assemble(blocks)


def weekly_text(iso_label, week, df, s, prev_avg, liq, wk):
    total = df["val"].sum()
    head = [f"📊 МХБ 7 хоног {iso_label} ({week[0][5:]} → {week[-1][5:]}) · {len(df):,} хэлцэл · {mn(total)} · "
            f"{len(week)} арилжааны өдөр · {len(s)} компани"]
    if prev_avg:
        head.append(f"Өдрийн дундаж арилжааны дүн өмнөх 7 хоногоос {pct((total / len(week) / prev_avg - 1) * 100)}")
    days = ["📅 Өдөр бүр (арилжааны дүн · хэлцэл · МХБ-н 20 тэнцүү/жигнэсэн өөрчлөлт):"]
    for _, r in wk["days"].iterrows():
        idx = f"{r['eq']:+.2f}% / {r['vw']:+.2f}%" if pd.notna(r["eq"]) else "—"
        days.append(f"{r['Гараг'][:2]} {sm(r['val'])} · {int(r['deals']):,} · {idx}")
    streak = []
    if wk["n_ret"] >= 3:
        ups = wk["updays"]
        allup = [k for k in s["Symbol"] if ups.get(k, (0, 0))[1] >= 3 and ups[k][0] == ups[k][1]][:4]
        alldn = [k for k in s["Symbol"] if ups.get(k, (0, 0))[1] >= 3 and ups[k][0] == 0][:4]
        if allup:
            streak.append("📶 Бүх өдөр өссөн: " + ", ".join(allup))
        if alldn:
            streak.append("📶 Бүх өдөр унасан: " + ", ".join(alldn))
    t = s[s["Хэлцэл"] >= MIN_DEALS]
    vol = t.sort_values("Өдрийн дундаж хэлбэлзэл (%)", ascending=False)
    calm = vol.tail(3).iloc[::-1]
    volb = ["🌊 Хамгийн хэлбэлзэлтэй: " + ", ".join(f"{r['Symbol']} {r['Өдрийн дундаж хэлбэлзэл (%)']:.1f}%/өдр" for _, r in vol.head(3).iterrows()),
            "🧘 Хамгийн тайван: " + ", ".join(f"{r['Symbol']} {r['Өдрийн дундаж хэлбэлзэл (%)']:.1f}%/өдр" for _, r in calm.iterrows())]
    corr = []
    if wk["corr"] is not None:
        pairs = wk["corr"]
        pos = pairs[pairs["r"] > 0].head(2)
        neg = pairs[pairs["r"] < 0].tail(2).iloc[::-1]
        if len(pos):
            corr.append("🔗 Хамт хөдөлсөн: " + ", ".join(f"{a}–{b} {r:+.2f}" for a, b, r in pos[["a", "b", "r"]].itertuples(index=False)))
        if len(neg):
            corr.append("🔗 Эсрэг хөдөлсөн: " + ", ".join(f"{a}–{b} {r:+.2f}" for a, b, r in neg[["a", "b", "r"]].itertuples(index=False)))
    blocks = [
        blk(0, head),
        blk(1, days),
        blk(1, index_block(s, weekly=True)),
        blk(2, fn=lambda n: price_block(s, weekly=True, limit=n), n=len(s), drop=False),
        blk(2, movers_block(s)),
        blk(2, close_block(s)),
        blk(3, streak),
        blk(2, fn=lambda n: biggest_block(df, weekly=True, limit=n), n=len(s), drop=False),
        blk(1, liquidity_block(s, "7 хоногийн", liq)),
        blk(4, volb),
        blk(5, corr),
        blk(4, hours_block(df, with_open=False)),
        blk(0, [DISCLAIMER]),
    ]
    return assemble(blocks)


def write_text(path, latest_name, text):
    path.write_text(text + "\n", encoding="utf-8")
    shutil.copyfile(path, REPORT_DIR / latest_name)


# ---------------------------------------------------------------- өдөр
def run_daily(arg=None):
    days = all_days()
    if not days:
        print("data/ хавтсанд өгөгдөл алга.")
        return 1
    day = arg or days[-1]
    if day not in days:
        print(f"{day} өдрийн файл олдсонгүй.")
        return 1
    df = load_day(day)
    if df.empty:
        print("Хэлцэл алга.")
        return 1
    earlier = [d for d in days if d < day]
    s = summarize(df, day_stats(earlier[-1]) if earlier else None)
    liq, prev_info = None, None
    if earlier:
        pdf = load_day(earlier[-1])
        base = pdf.groupby("sym")["val"].sum().to_dict()
        cur = df.groupby("sym")["val"].sum().to_dict()
        liq = (cur, base, "өмнөх өдрөөс")
        prev_info = (earlier[-1], float(pdf["val"].sum()), len(pdf))
    sheets = {
        "Тойм": s,
        "Цагаар": hourly(df),
        "Том хэлцэл (компаниар)": biggest_table(df),
        "Том хэлцэл": block_table(df),
        "Үнийн үсрэлт": jump_table(df),
        "Нөхцөлтэй хэлцэл": condition_table(df),
        "Тайлбар": definitions(),
    }
    if liq:
        sheets = {**{k: v for k, v in sheets.items() if k != "Тайлбар"},
                  "Хөрвөх чадварын өөрчлөлт": liq_table(liq[0], liq[1]), "Тайлбар": sheets["Тайлбар"]}
    REPORT_DIR.mkdir(parents=True, exist_ok=True)
    save_book(REPORT_DIR / f"daily_{day}.xlsx", sheets)
    write_text(REPORT_DIR / f"daily_{day}.txt", "latest_daily.txt", daily_text(day, df, s, liq, prev_info))
    print(f"Өдрийн шинжилгээ бэлэн: {REPORT_DIR}/daily_{day}.xlsx")
    return 0


# ---------------------------------------------------------------- 7 хоног
def prev_for_week(days, monday):
    before = [d for d in days if date.fromisoformat(d) < monday]
    if not before:
        return None
    prev = day_stats(before[-1])                     # өмнөх хаалт: өмнөх арилжааны өдрийнх
    pw = [d for d in before if date.fromisoformat(d) >= monday - timedelta(days=7)]
    if pw:                                           # өмнөх VWAP: өмнөх 7 хоногийн нийт арилжаанаас
        pdf = pd.concat([load_day(d) for d in pw], ignore_index=True)
        for sym, g in pdf.groupby("sym"):
            prev.setdefault(sym, {})["vwap"] = g["val"].sum() / g["qty"].sum()
    return prev


def run_weekly(arg=None):
    days = all_days()
    if not days:
        print("data/ хавтсанд өгөгдөл алга.")
        return 1
    ref = date.fromisoformat(arg or days[-1])
    monday = ref - timedelta(days=ref.weekday())
    week = [d for d in days if monday <= date.fromisoformat(d) <= monday + timedelta(days=6)]
    if not week:
        print("Тухайн 7 хоногт өгөгдөл алга.")
        return 1
    df = pd.concat([load_day(d) for d in week], ignore_index=True).sort_values(["dt", "id"]).reset_index(drop=True)
    prev = prev_for_week(days, monday)
    s = summarize(df, prev)

    grp = df.groupby(["sym", "day"])
    dclose = grp["px"].last().unstack("day")
    dval = grp["val"].sum().unstack("day").fillna(0.0)
    rng = grp["px"].agg(["max", "min"])
    rng["r"] = (rng["max"] / rng["min"] - 1) * 100
    ffilled = dclose.ffill(axis=1)
    ret = (ffilled / ffilled.shift(1, axis=1) - 1) * 100

    s["Арилжсан өдөр"] = s["Symbol"].map(dclose.notna().sum(axis=1)).astype(int)
    s["Өдрийн дундаж эргэлт (₮)"] = s["Эргэлт (₮)"] / len(week)
    s["Өдрийн дундаж хэлбэлзэл (%)"] = s["Symbol"].map(rng.groupby("sym")["r"].mean())
    s["Өдрийн өгөөжийн σ (%)"] = s["Symbol"].map(ret.std(axis=1))
    s["Хамгийн идэвхтэй өдөр"] = s["Symbol"].map(dval.idxmax(axis=1)).map(wd)

    mkt = df.groupby("day").agg(deals=("id", "count"), qty=("qty", "sum"), val=("val", "sum")).reset_index()
    mkt["Гараг"] = mkt["day"].map(wd)
    mkt["Өмнөх өдрөөс (%)"] = (mkt["val"] / mkt["val"].shift(1) - 1) * 100
    mkt = mkt.rename(columns={"day": "Огноо", "deals": "Хэлцэл", "qty": "Ширхэг", "val": "Эргэлт (₮)"})[
        ["Огноо", "Гараг", "Хэлцэл", "Ширхэг", "Эргэлт (₮)", "Өмнөх өдрөөс (%)"]]

    turn = dval.copy()
    turn["Нийт"] = turn.sum(axis=1)
    turn = turn.sort_values("Нийт", ascending=False).rename_axis(index="Symbol", columns=None).reset_index()
    close_tbl = dclose.rename_axis(index="Symbol", columns=None).reset_index()

    sheets = {"Тойм": s, "Зах зээл өдрөөр": mkt, "Эргэлт өдрөөр": turn, "Өдрийн сүүлийн хэлцэл": close_tbl}
    if len(week) >= 4:  # 3+ өдрийн өгөөж байж хамаарал утга агуулна
        corr = ret.T.corr(min_periods=3).round(2).rename_axis(index="Symbol", columns=None).reset_index()
        sheets["Хамаарал"] = corr
    sheets.update({"Том хэлцэл (компаниар)": biggest_table(df), "Том хэлцэл": block_table(df), "Үнийн үсрэлт": jump_table(df),
                   "Цагаар": hourly(df), "Нөхцөлтэй хэлцэл": condition_table(df), "Тайлбар": definitions()})

    iso = monday.isocalendar()
    label = f"{iso.year}-W{iso.week:02d}"
    prev_days = [d for d in days if monday - timedelta(days=7) <= date.fromisoformat(d) < monday]
    prev_avg = sum(load_day(d)["val"].sum() for d in prev_days) / len(prev_days) if prev_days else None

    # өдөр бүрийн индекс, дараалсан өсөлт/уналт, хамаарал (хаалт = өдрийн жигнэсэн дундаж ханш)
    g_val, g_qty = df.groupby(["sym", "day"])["val"].sum(), df.groupby(["sym", "day"])["qty"].sum()
    dvw, dv = (g_val / g_qty).unstack("day"), g_val.unstack("day").fillna(0.0)
    prevc = pd.Series({k: v["close"] for k, v in (prev or {}).items()}, dtype=float).reindex(dvw.index) if prev else \
        pd.Series(float("nan"), index=dvw.index)
    rch = dvw.div(dvw.ffill(axis=1).shift(1, axis=1)) - 1
    rch[dvw.columns[0]] = dvw[dvw.columns[0]] / prevc - 1
    rows_ = []
    for d_ in dvw.columns:
        c_, m_ = rch[d_], rch[d_].notna()
        w_ = dv[d_][m_]
        rows_.append({"day": d_, "eq": c_[m_].mean() * 100 if m_.any() else float("nan"),
                      "vw": (c_[m_] * w_).sum() / w_.sum() * 100 if w_.sum() > 0 else float("nan")})
    day_idx = pd.DataFrame(rows_)
    days_tbl = mkt.rename(columns={"Огноо": "day", "Хэлцэл": "deals", "Эргэлт (₮)": "val"})[["day", "Гараг", "deals", "val"]].merge(day_idx, on="day")
    mkt["МХБ-н 20 тэнцүү (%)"] = days_tbl["eq"].values
    mkt["МХБ-н 20 жигнэсэн (%)"] = days_tbl["vw"].values
    n_ret = int(rch.notna().any(axis=0).sum())
    updays = {k: (int((rch.loc[k] > 0).sum()), int(rch.loc[k].notna().sum())) for k in rch.index}
    corr_pairs = None
    if n_ret >= 3:
        cm = rch.T.corr(min_periods=3)
        cp = [(a_, b_, cm.loc[a_, b_]) for i_, a_ in enumerate(cm.index) for b_ in cm.index[i_ + 1:] if pd.notna(cm.loc[a_, b_])]
        corr_pairs = pd.DataFrame(cp, columns=["a", "b", "r"]).sort_values("r", ascending=False).reset_index(drop=True)
    wk = {"days": days_tbl, "n_ret": n_ret, "updays": updays, "corr": corr_pairs}

    liq = None
    if prev_days:
        base = pd.concat([load_day(d) for d in prev_days]).groupby("sym")["val"].sum() / len(prev_days)
        cur = df.groupby("sym")["val"].sum() / len(week)
        liq = (cur.to_dict(), base.to_dict(), "өмнөх 7 хоногоос, өдрийн дундажаар")
    elif len(week) >= 2:  # өмнөх 7 хоногийн өгөгдөл байхгүй: сүүлийн өдрийг 7 хоногийн өмнөх өдрүүдтэй харьцуулна
        last = week[-1]
        base = df[df["day"] != last].groupby("sym")["val"].sum() / (len(week) - 1)
        cur = df[df["day"] == last].groupby("sym")["val"].sum()
        liq = (cur.to_dict(), base.to_dict(),
               f"{wd(last)[:2]} гарагийн дүн vs өмнөх өдрүүдийн дундаж")
    if liq:
        sheets = {**{k: v for k, v in sheets.items() if k != "Тайлбар"},
                  "Хөрвөх чадварын өөрчлөлт": liq_table(liq[0], liq[1]), "Тайлбар": sheets["Тайлбар"]}
    REPORT_DIR.mkdir(parents=True, exist_ok=True)
    save_book(REPORT_DIR / f"weekly_{label}.xlsx", sheets)
    write_text(REPORT_DIR / f"weekly_{label}.txt", "latest_weekly.txt",
               weekly_text(label, week, df, s, prev_avg, liq, wk))
    print(f"7 хоногийн шинжилгээ бэлэн: {REPORT_DIR}/weekly_{label}.xlsx ({len(week)} өдөр)")
    return 0


def ub_today():
    if os.environ.get("TODAY"):  # туршилтад
        return date.fromisoformat(os.environ["TODAY"])
    return (datetime.now(timezone.utc) + timedelta(hours=8)).date()  # Улаанбаатар UTC+8


def week_path(monday):
    iso = monday.isocalendar()
    return REPORT_DIR / f"weekly_{iso.year}-W{iso.week:02d}.xlsx"


def run_weekly_auto():
    """Баасан гарагт энэ 7 хоногийг гаргана. Бусад өдөр өмнөх 7 хоногийн тайлан
    хараахан гараагүй бол (жишээ нь Баасан гарагийн ажил алдагдсан) нөхөж гаргана.
    Бирж хэдэн өдөр ажиллаагүй ч байгаа өдрүүдээр тайлан гарна."""
    today = ub_today()
    monday = today - timedelta(days=today.weekday())
    if today.weekday() == 4:
        target = monday
    else:
        target = monday - timedelta(days=7)
        if week_path(target).exists():
            print("Өмнөх 7 хоногийн тайлан гарсан тул алгасав.")
            return 0
    end = target + timedelta(days=6)
    if not [d for d in all_days() if target <= date.fromisoformat(d) <= end]:
        print(f"{target} долоо хоногт арилжааны өгөгдөл алга — тайлан гаргахгүй.")
        return 0
    return run_weekly(target.isoformat())


if __name__ == "__main__":
    mode = sys.argv[1] if len(sys.argv) > 1 else "daily"
    arg = sys.argv[2] if len(sys.argv) > 2 else None
    if mode == "daily":
        sys.exit(run_daily(arg))
    if mode == "weekly":
        sys.exit(run_weekly(arg))
    if mode == "weekly-auto":
        sys.exit(run_weekly_auto())
    print(__doc__)
    sys.exit(2)
